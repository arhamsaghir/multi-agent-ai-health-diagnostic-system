"""Generate PowerPoint presentation for B.Tech Final Year Project.

Multi-Agent AI Health Diagnostic System
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ─── Color Palette ─────────────────────────────────────────────
DARK_BG = RGBColor(0x0F, 0x17, 0x2A)       # Dark navy
ACCENT_BLUE = RGBColor(0x3B, 0x82, 0xF6)   # Bright blue
ACCENT_GREEN = RGBColor(0x10, 0xB9, 0x81)  # Emerald green
ACCENT_ORANGE = RGBColor(0xF5, 0x9E, 0x0B) # Amber
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0x94, 0xA3, 0xB8)
DARK_GRAY = RGBColor(0x33, 0x41, 0x55)
TABLE_HEADER = RGBColor(0x1E, 0x3A, 0x5F)
TABLE_ROW = RGBColor(0x1E, 0x29, 0x3B)
TABLE_ALT = RGBColor(0x17, 0x20, 0x33)


def set_slide_bg(slide, color):
    """Set solid background color for a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 bold=False, color=WHITE, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    """Add a text box to a slide."""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16,
                    color=LIGHT_GRAY, spacing=6, bullet_char="•"):
    """Add a bullet-point list to a slide."""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"{bullet_char} {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(spacing)
    return txBox


def add_accent_line(slide, left, top, width, color=ACCENT_BLUE, thickness=3):
    """Add a horizontal accent line."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(thickness / 72)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_card(slide, left, top, width, height, fill_color=DARK_GRAY):
    """Add a rounded card shape."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return shape


def create_presentation():
    """Create the full presentation."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ════════════════════════════════════════════════════════════
    # SLIDE 1: Title Slide
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, DARK_BG)

    add_accent_line(slide, 1.5, 2.2, 10.3, ACCENT_BLUE, 4)

    add_text_box(slide, 1.5, 2.5, 10.3, 1.5,
                 "Multi-Agent AI Health Diagnostic System",
                 font_size=40, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, 1.5, 4.0, 10.3, 0.8,
                 "B.Tech Final Year Project — Computer Science & Engineering",
                 font_size=22, color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, 1.5, 4.8, 10.3, 0.5,
                 "Presented by: [Your Name]  |  Roll No: [Your Roll Number]",
                 font_size=18, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, 1.5, 5.3, 10.3, 0.5,
                 "Under the Guidance of: [Supervisor Name]",
                 font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    add_accent_line(slide, 1.5, 6.0, 10.3, ACCENT_BLUE, 4)

    # Speaker notes
    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "Welcome. Today I present my final year project: a Multi-Agent AI Health Diagnostic System.\n"
        "This system uses 8 specialized AI agents working together — similar to how a medical team operates —\n"
        "to analyze patient symptoms, medical images, and clinical knowledge for diagnostic reports.\n"
        "Let me walk you through the problem, architecture, implementation, and results."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 2: Problem Statement
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "Problem Statement", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_BLUE, 3)

    add_text_box(slide, 1.0, 1.6, 11.3, 1.0,
                 "Current AI diagnostic systems have critical limitations:",
                 font_size=22, color=ACCENT_ORANGE)

    items = [
        "Single-modality: Analyze either text OR images, not both simultaneously",
        "Black-box models: No explanation of reasoning process",
        "No evidence grounding: Predictions without clinical reference",
        "Monolithic architecture: Hard to maintain, update, or extend",
        "Poor uncertainty handling: Cannot express confidence across dimensions",
        "Healthcare access gap: Limited specialist availability in underserved areas"
    ]
    add_bullet_list(slide, 1.5, 2.8, 10.3, 4.0, items, font_size=20, spacing=12)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "The core problem is that current AI diagnostic tools are limited.\n"
        "They process one type of data at a time — either text or images.\n"
        "They don't explain their reasoning, and they don't reference actual clinical guidelines.\n"
        "Meanwhile, healthcare faces an access crisis — not enough specialists in many regions.\n"
        "Our system addresses all these gaps simultaneously."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 3: Objectives
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "Project Objectives", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_GREEN, 3)

    objectives = [
        "Design a multi-agent AI architecture emulating clinical team reasoning",
        "Develop specialized agents for text, images, knowledge, and risk prediction",
        "Create an orchestration layer coordinating outputs into unified reports",
        "Build a web UI with voice input, demo mode, and report export",
        "Evaluate diagnostic accuracy, response time, and usability"
    ]
    for i, obj in enumerate(objectives):
        y = 1.8 + i * 0.9
        card = add_card(slide, 1.5, y, 10.3, 0.7, DARK_GRAY)
        add_text_box(slide, 2.0, y + 0.08, 9.5, 0.55,
                     f"{i+1}. {obj}", font_size=18, color=WHITE)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "Our five objectives were clear:\n"
        "1) Design the multi-agent architecture\n"
        "2) Build each specialized agent\n"
        "3) Create the orchestration layer\n"
        "4) Build the user interface\n"
        "5) Evaluate the system\n"
        "Each objective maps to a chapter of the implementation."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 4: System Architecture
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "System Architecture", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_BLUE, 3)

    # Three-tier architecture boxes
    tiers = [
        ("Frontend Tier", "Streamlit Web App\nDiagnosis | Vision | Chat | History", ACCENT_BLUE),
        ("API Tier", "FastAPI REST API\n/api/diagnose | /api/vision | /api/chat", ACCENT_GREEN),
        ("Agent Tier", "8 Specialized Agents\nNLP | Vision | RAG | Prediction | Voice | Orchestrator | Synthesis | Scorer", ACCENT_ORANGE),
    ]

    for i, (title, desc, color) in enumerate(tiers):
        y = 1.8 + i * 1.6
        card = add_card(slide, 1.5, y, 10.3, 1.3, DARK_GRAY)
        add_text_box(slide, 1.8, y + 0.05, 3, 0.4, title, font_size=16, bold=True, color=color)
        add_text_box(slide, 1.8, y + 0.45, 9.5, 0.8, desc, font_size=15, color=LIGHT_GRAY)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "The system follows a three-tier architecture.\n"
        "Top layer: Streamlit frontend with 5 pages for different use cases.\n"
        "Middle layer: FastAPI REST API handling all requests.\n"
        "Bottom layer: 8 specialized AI agents orchestrated together.\n"
        "Data flows from the patient input through multiple agents, gets scored for confidence,\n"
        "then synthesized into a final diagnostic report."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 5: The 8 Agents
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "The 8 Specialized AI Agents", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_ORANGE, 3)

    agents = [
        ("NLP Agent", "Symptom extraction & body system detection", "BioBERT / Regex"),
        ("Vision Agent", "X-ray classification (15 diseases)", "ResNet-18 / PyTorch"),
        ("RAG Agent", "Clinical knowledge retrieval", "FAISS + LLM"),
        ("Prediction Agent", "Disease risk scoring from vitals", "GradientBoosting"),
        ("Voice Agent", "Speech-to-text transcription", "OpenAI Whisper"),
        ("Orchestrator", "Agent coordination & task delegation", "LangGraph"),
        ("Synthesis Agent", "Diagnostic report generation", "OpenRouter LLM"),
        ("Confidence Scorer", "Multi-agent confidence aggregation", "Weighted Ensemble"),
    ]

    for i, (name, desc, tech) in enumerate(agents):
        col = i % 4
        row = i // 4
        x = 1.0 + col * 3.0
        y = 1.8 + row * 2.8

        card = add_card(slide, x, y, 2.7, 2.4, DARK_GRAY)
        add_text_box(slide, x + 0.15, y + 0.1, 2.4, 0.5,
                     name, font_size=15, bold=True, color=ACCENT_BLUE)
        add_accent_line(slide, x + 0.15, y + 0.6, 2.4, ACCENT_BLUE, 1)
        add_text_box(slide, x + 0.15, y + 0.7, 2.4, 0.9,
                     desc, font_size=12, color=LIGHT_GRAY)
        add_text_box(slide, x + 0.15, y + 1.6, 2.4, 0.6,
                     f"Tech: {tech}", font_size=11, color=ACCENT_ORANGE)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "Each agent specializes in one diagnostic modality:\n"
        "NLP extracts symptoms from text. Vision analyzes X-rays. RAG retrieves medical knowledge.\n"
        "Prediction calculates risk scores. Voice transcribes speech.\n"
        "The Orchestrator coordinates everything. Synthesis generates the final report.\n"
        "And the Confidence Scorer aggregates all outputs with weighted confidence.\n"
        "This mirrors how a real medical team operates."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 6: Data Flow
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "Data Flow Pipeline", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_GREEN, 3)

    steps = [
        ("1. Patient Input", "Text description or voice recording", ACCENT_BLUE),
        ("2. NLP Processing", "Extracts symptoms, body systems, severity, duration", ACCENT_BLUE),
        ("3. Parallel Analysis", "Vision Agent | RAG Agent | Prediction Agent", ACCENT_GREEN),
        ("4. Confidence Scoring", "Weighted ensemble: NLP 30%, Vision 25%, RAG 25%, Prediction 20%", ACCENT_ORANGE),
        ("5. Synthesis", "LLM generates formatted diagnostic report", ACCENT_ORANGE),
        ("6. Output", "Report with conditions, severity, recommendations, red flags", ACCENT_GREEN),
    ]

    for i, (title, desc, color) in enumerate(steps):
        y = 1.8 + i * 0.85
        card = add_card(slide, 1.5, y, 10.3, 0.7, DARK_GRAY)
        add_text_box(slide, 1.8, y + 0.05, 3, 0.55, title, font_size=16, bold=True, color=color)
        add_text_box(slide, 5.0, y + 0.05, 6.5, 0.55, desc, font_size=14, color=LIGHT_GRAY)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "The data flow is sequential but with parallel processing in the middle.\n"
        "First, patient input goes through NLP to extract structured symptoms.\n"
        "Then three agents run in parallel — Vision for images, RAG for knowledge, Prediction for vitals.\n"
        "Their outputs are combined by the Confidence Scorer using weighted ensemble.\n"
        "Finally, the Synthesis Agent generates a human-readable diagnostic report.\n"
        "The entire pipeline takes 10-30 seconds with live API, under 1 second in demo mode."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 7: Technology Stack
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "Technology Stack", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_BLUE, 3)

    tech_table = [
        ("Layer", "Technology", "Purpose"),
        ("Frontend", "Streamlit", "Rapid web UI with multi-page support"),
        ("API", "FastAPI", "Async REST API with auto-generated docs"),
        ("NLP", "Hugging Face Transformers", "BioBERT for biomedical text"),
        ("Vision", "PyTorch + TorchVision", "ResNet-18 for image classification"),
        ("ML", "scikit-learn", "GradientBoosting for risk prediction"),
        ("Vector Store", "FAISS", "Efficient similarity search for RAG"),
        ("Orchestration", "LangGraph", "Multi-agent state machine"),
        ("LLM", "OpenRouter API", "Free-tier access to GPT-4-class models"),
        ("Voice", "OpenAI Whisper", "Industry-leading speech recognition"),
        ("Database", "SQLite", "Zero-config local storage"),
    ]

    rows = len(tech_table)
    cols = 3
    table_shape = slide.shapes.add_table(rows, cols, Inches(1.0), Inches(1.6), Inches(11.3), Inches(5.2))
    table = table_shape.table

    col_widths = [Inches(2.5), Inches(2.8), Inches(6.0)]
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = tech_table[r][c]
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(13)
            p.font.name = "Calibri"
            if r == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
            else:
                p.font.color.rgb = LIGHT_GRAY
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "The technology stack was chosen for reliability, community support, and cost-effectiveness.\n"
        "All LLM operations use OpenRouter's free tier — zero cost.\n"
        "Streamlit enables rapid UI prototyping. FastAPI provides async API.\n"
        "We replaced XGBoost with scikit-learn due to Python 3.13 compatibility issues.\n"
        "Everything runs locally with SQLite — no cloud dependencies for core functionality."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 8: Results
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "Results & Performance", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_GREEN, 3)

    # Metrics cards
    metrics = [
        ("Vision Model", "Val Loss: 0.0562", "4,500 images, 15 classes"),
        ("NLP Agent", "200+ symptom terms", "11 body systems covered"),
        ("Test Suite", "70 tests passing", "52 unit + 18 API"),
        ("Response Time", "< 1s (Demo)", "10-30s (Live API)"),
        ("Dataset", "500 patient records", "Synthetic with real clinical ranges"),
        ("Code Base", "8,000+ lines", "50+ Python files"),
    ]

    for i, (title, metric, detail) in enumerate(metrics):
        col = i % 3
        row = i // 3
        x = 1.0 + col * 4.0
        y = 1.8 + row * 2.6

        card = add_card(slide, x, y, 3.6, 2.2, DARK_GRAY)
        add_text_box(slide, x + 0.2, y + 0.1, 3.2, 0.4, title, font_size=15, bold=True, color=ACCENT_GREEN)
        add_accent_line(slide, x + 0.2, y + 0.55, 3.2, ACCENT_GREEN, 1)
        add_text_box(slide, x + 0.2, y + 0.7, 3.2, 0.6, metric, font_size=22, bold=True, color=WHITE)
        add_text_box(slide, x + 0.2, y + 1.4, 3.2, 0.6, detail, font_size=13, color=LIGHT_GRAY)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "Key results:\n"
        "Vision model achieved 0.0562 validation loss on 4,500 synthetic X-rays.\n"
        "NLP agent covers 200+ symptom terms across 11 body systems.\n"
        "70 automated tests all passing — 52 unit tests and 18 API tests.\n"
        "Demo mode provides instant results under 1 second.\n"
        "Live API takes 10-30 seconds on CPU.\n"
        "The codebase spans over 8,000 lines across 50+ files."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 9: User Interface
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "User Interface Features", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_BLUE, 3)

    features = [
        "5-page Streamlit app: Diagnosis, Vision, Chat, History, About",
        "Demo Mode toggle for instant results without API server",
        "Voice input via browser microphone (Whisper transcription)",
        "Severity badges: color-coded green/orange/red",
        "Confidence bars for each detected condition",
        "Red flag detection for emergency indicators",
        "Report export: PDF, CSV, and Text formats",
        "History page with SQLite persistence, filters, and search",
        "Responsive design with custom CSS theming"
    ]
    add_bullet_list(slide, 1.5, 2.0, 10.3, 5.0, features, font_size=20, spacing=14)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "The user interface was designed for accessibility and ease of demonstration.\n"
        "Five pages cover all functionality. Demo Mode is crucial for smooth presentations.\n"
        "Voice input makes the system accessible. History persists across browser sessions.\n"
        "Report export in multiple formats supports sharing with healthcare providers.\n"
        "Color-coded severity provides immediate visual understanding of urgency."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 10: Challenges & Solutions
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "Challenges & Solutions", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_ORANGE, 3)

    challenges = [
        ("XGBoost Python 3.13 Incompatibility",
         "Segfault during model deserialization",
         "Replaced with scikit-learn GradientBoosting"),
        ("No Real Medical Datasets",
         "Limited access to labeled clinical images",
         "Generated 4,500 synthetic X-rays with disease patterns"),
        ("CPU-Only Inference Speed",
         "Model inference too slow for live demos",
         "Built Demo Mode with pre-computed results"),
        ("Multi-Agent Orchestration",
         "Coordinating 8 agents without race conditions",
         "LangGraph state machine with sequential + parallel execution"),
    ]

    for i, (challenge, problem, solution) in enumerate(challenges):
        y = 1.8 + i * 1.3
        card = add_card(slide, 1.5, y, 10.3, 1.1, DARK_GRAY)
        add_text_box(slide, 1.8, y + 0.05, 3, 0.4, challenge, font_size=14, bold=True, color=ACCENT_ORANGE)
        add_text_box(slide, 5.0, y + 0.05, 6.5, 0.4, f"Problem: {problem}", font_size=12, color=LIGHT_GRAY)
        add_text_box(slide, 5.0, y + 0.5, 6.5, 0.4, f"Solution: {solution}", font_size=12, color=ACCENT_GREEN)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "We faced four major challenges:\n"
        "1) XGBoost crashed on Python 3.13 — solved by switching to scikit-learn.\n"
        "2) No real medical data — generated synthetic datasets with realistic patterns.\n"
        "3) Slow CPU inference — built Demo Mode for instant results during presentations.\n"
        "4) Complex agent coordination — used LangGraph's state machine for reliable orchestration.\n"
        "Each challenge taught us important lessons about production AI development."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 11: Future Work
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_text_box(slide, 1.0, 0.5, 11.3, 0.8,
                 "Future Work", font_size=36, bold=True, color=WHITE)
    add_accent_line(slide, 1.0, 1.3, 3, ACCENT_GREEN, 3)

    future = [
        "Train on real clinical datasets (NIH ChestX-ray14, MIMIC-IV)",
        "GPU deployment for real-time inference",
        "Expand knowledge base to 100+ clinical guidelines",
        "Clinical validation against physician assessments",
        "FDA/CE regulatory pathway exploration (SaMD classification)",
        "Multilingual support with XLM-RoBERTa",
        "Wearable device integration for continuous monitoring",
        "Federated learning for multi-institution training",
        "SHAP/LIME explainability visualizations for clinicians",
        "Mobile app with on-device inference"
    ]
    add_bullet_list(slide, 1.5, 1.8, 10.3, 5.5, future, font_size=18, spacing=12)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "The roadmap for future work includes:\n"
        "Real clinical data for validation. GPU deployment for speed.\n"
        "Regulatory pathway exploration. Multilingual support.\n"
        "Federated learning for privacy-preserving multi-institution training.\n"
        "Explainability tools so clinicians can trust the AI's reasoning.\n"
        "And eventually, a mobile app for point-of-care use."
    )

    # ════════════════════════════════════════════════════════════
    # SLIDE 12: Conclusion & Thank You
    # ════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_accent_line(slide, 1.5, 2.0, 10.3, ACCENT_BLUE, 4)

    add_text_box(slide, 1.5, 2.3, 10.3, 1.0,
                 "Conclusion", font_size=40, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, 1.5, 3.5, 10.3, 2.0,
                 "This project demonstrates that multi-agent AI architectures can provide "
                 "structured, interpretable, and evidence-based diagnostic assistance "
                 "while maintaining computational efficiency suitable for deployment on standard hardware.",
                 font_size=20, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, 1.5, 5.5, 10.3, 0.6,
                 "Thank You", font_size=36, bold=True, color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, 1.5, 6.2, 10.3, 0.5,
                 "Questions & Discussion", font_size=20, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    add_accent_line(slide, 1.5, 6.8, 10.3, ACCENT_BLUE, 4)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "In conclusion, this project successfully demonstrates a multi-agent AI health diagnostic system.\n"
        "The system integrates 8 specialized agents, processes multiple data modalities,\n"
        "and generates evidence-based diagnostic reports.\n"
        "The modular architecture enables independent updates and future expansion.\n"
        "Thank you for your attention. I'm happy to take questions."
    )

    # Save
    output_path = "/Users/arhamsaghir/Documents/Final Year Project/docs/presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    create_presentation()
